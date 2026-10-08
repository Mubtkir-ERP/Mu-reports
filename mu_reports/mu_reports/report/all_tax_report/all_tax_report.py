# Copyright (c) 2026, Mubtkir and contributors
from collections import defaultdict
import json
import math
import frappe
from frappe import _
from frappe.utils import flt, cint, getdate, now_datetime
from frappe.query_builder import DocType

SECTION_NAMES = {
    'Sales Invoice': _('Sales Invoices'), 'Sales Return': _('Sales Returns'),
    'Purchase Invoice': _('Purchase Invoices'), 'Purchase Return': _('Purchase Returns'),
    'Payment Entry': _('ERPNext Payment Entries'), 'Vouchers Entry': _('Voucher Entries'),
    'Journal Entry': _('Journal Entries'), 'Other': _('Other GL Vouchers'),
}
SECTION_ORDER = tuple(SECTION_NAMES)
RETURN_LABELS = {
    1: 'المبيعات الخاضعة للنسبة الأساسية',
    2: 'المبيعات للمواطنين: الخدمات الصحية والتعليم الأهلي',
    3: 'المبيعات المحلية الخاضعة للنسبة الصفرية', 4: 'الصادرات', 5: 'المبيعات المعفاة',
    7: 'المشتريات الخاضعة للنسبة الأساسية',
    8: 'الاستيرادات الخاضعة للنسبة الأساسية والمدفوعة عند الاستيراد',
    9: 'الاستيرادات الخاضعة للنسبة الأساسية وفق الاحتساب العكسي',
    10: 'المشتريات الخاضعة للنسبة الصفرية', 11: 'المشتريات المعفاة',
}


def _json(value, default):
    if not value:
        return default
    try:
        return json.loads(value) if isinstance(value, str) else value
    except (ValueError, TypeError):
        frappe.throw(_('Invalid classification settings.'))


def normalize_tax_accounts(value):
    if isinstance(value, str):
        try:
            value = json.loads(value)
        except ValueError:
            value = value.replace(',', '\n').splitlines()
    if isinstance(value, str):
        value = [value]
    return list(dict.fromkeys(str(v.get('value') if isinstance(v, dict) else v).strip()
                              for v in (value or []) if v and (not isinstance(v, dict) or v.get('value'))))


def _settings(value):
    rows = _json(value, [])
    if not isinstance(rows, list) or any(not isinstance(r, dict) for r in rows):
        frappe.throw(_('Classification settings must contain rows.'))
    return rows


def _validate(filters):
    from frappe.desk.query_report import get_report_doc
    get_report_doc('All Tax Report')
    if not filters.get('company'):
        frappe.throw(_('Please select a Company.'))
    if not frappe.has_permission('Company', 'read', filters.company):
        frappe.throw(_('Not permitted to read this Company.'), frappe.PermissionError)
    if not filters.get('from_date') or not filters.get('to_date'):
        frappe.throw(_('Please select both dates.'))
    if getdate(filters.from_date) > getdate(filters.to_date):
        frappe.throw(_('From Date must not be after To Date.'))
    accounts = normalize_tax_accounts(filters.get('tax_accounts'))
    if not accounts:
        frappe.throw(_('Please select at least one Tax Account to filter by.'))
    valid = frappe.get_all('Account', filters={'name': ['in', accounts], 'company': filters.company,
                                              'is_group': 0}, pluck='name')
    if set(valid) != set(accounts):
        frappe.throw(_('Selected accounts must be ledger accounts of the selected Company.'))
    filters.tax_accounts = accounts
    return filters


def _field_exists(doctype, field):
    return frappe.get_meta(doctype).has_field(field)


def _master_fields(doctype, names, fields):
    if not names or not frappe.db.exists('DocType', doctype):
        return {}
    fields = ['name'] + [f for f in fields if _field_exists(doctype, f)]
    return {r.name: r for r in frappe.get_all(doctype, filters={'name': ['in', list(names)]}, fields=fields)}


def _section(vtype, doc):
    if vtype in ('Sales Invoice', 'Purchase Invoice') and doc and doc.get('is_return'):
        return 'Sales Return' if vtype == 'Sales Invoice' else 'Purchase Return'
    return vtype if vtype in SECTION_NAMES else 'Other'


def _invoice_net(vtype, doc):
    if vtype not in ('Sales Invoice', 'Purchase Invoice') or not doc or doc.get('base_net_total') is None:
        return None
    return flt(doc.base_net_total) * (-1 if vtype == 'Purchase Invoice' else 1)


def _read_ledger(filters):
    gle = DocType('GL Entry')
    query = (frappe.qb.from_(gle).select(gle.voucher_no, gle.voucher_type, gle.posting_date,
             gle.account, gle.party, gle.party_type, gle.debit, gle.credit)
             .where(gle.is_cancelled == 0).where(gle.company == filters.company)
             .where(gle.account.isin(filters.tax_accounts))
             .where(gle.posting_date >= filters.from_date).where(gle.posting_date <= filters.to_date))
    if filters.get('invoice_no'):
        query = query.where(gle.voucher_no == filters.invoice_no)
    return query.run(as_dict=True)


def _visible(row, filters):
    if filters.get('party') and row['party'] != filters.party:
        return False
    if filters.get('voucher_type') and row['voucher_type'] != filters.voucher_type:
        return False
    mode = filters.get('document_group') or 'All'
    section = row['section_key']
    if mode == 'Sales' and section not in ('Sales Invoice', 'Sales Return'):
        return False
    if mode == 'Purchases' and section not in ('Purchase Invoice', 'Purchase Return'):
        return False
    if mode == 'Returns' and section not in ('Sales Return', 'Purchase Return'):
        return False
    if mode == 'Settlements' and row['voucher_type'] in ('Sales Invoice', 'Purchase Invoice'):
        return False
    return cint(filters.get('include_zero_tax')) or abs(row['tax_amount']) > 0.0000001


def _load_documents(filters, ledger):
    vouchers = {}
    for r in ledger:
        key = (r.voucher_type, r.voucher_no)
        v = vouchers.setdefault(key, {'posting_date': r.posting_date, 'party': '', 'party_type': '',
                                      'accounts': {}, 'has_tax_gl': True})
        a = v['accounts'].setdefault(r.account, {'debit': 0.0, 'credit': 0.0})
        a['debit'] += flt(r.debit)
        a['credit'] += flt(r.credit)
        if r.party and not v['party']:
            v['party'], v['party_type'] = r.party, r.party_type or ''
    by_type = defaultdict(set)
    for vtype, name in vouchers:
        by_type[vtype].add(name)
    docs = {}
    for vtype, names in by_type.items():
        docs[vtype] = _master_fields(vtype, names, ['base_net_total', 'customer', 'supplier', 'party',
                                                  'party_type', 'is_return', 'payment_type', 'docstatus'])
        # Tax GL lines often have no party; recover it from the same voucher's GL.
        for r in frappe.get_all('GL Entry', filters={'company': filters.company, 'is_cancelled': 0,
                   'voucher_type': vtype, 'voucher_no': ['in', list(names)], 'party': ['!=', '']},
                   fields=['voucher_no', 'party', 'party_type'], order_by='name asc'):
            v = vouchers[(vtype, r.voucher_no)]
            if not v['party']:
                v['party'], v['party_type'] = r.party, r.party_type or ''
    if cint(filters.get('include_non_taxed')):
        for vtype in ('Sales Invoice', 'Purchase Invoice'):
            query_filters = {'company': filters.company, 'docstatus': 1,
                             'posting_date': ['between', [filters.from_date, filters.to_date]],
                             'total_taxes_and_charges': 0}
            if filters.get('invoice_no'):
                query_filters['name'] = filters.invoice_no
            fields = ['name', 'posting_date', 'base_net_total', 'is_return', 'docstatus',
                      'customer' if vtype == 'Sales Invoice' else 'supplier']
            for doc in frappe.get_all(vtype, filters=query_filters, fields=fields):
                key = (vtype, doc.name)
                if key not in vouchers:
                    vouchers[key] = {'posting_date': doc.posting_date, 'party': '', 'party_type': '',
                                     'accounts': {}, 'has_tax_gl': False}
                docs.setdefault(vtype, {})[doc.name] = doc
    customers, suppliers = set(), set()
    for (vtype, name), v in vouchers.items():
        doc = docs.get(vtype, {}).get(name) or {}
        if vtype == 'Sales Invoice':
            v['party'], v['party_type'] = doc.get('customer') or v['party'], 'Customer'
        elif vtype == 'Purchase Invoice':
            v['party'], v['party_type'] = doc.get('supplier') or v['party'], 'Supplier'
        else:
            v['party'] = v['party'] or doc.get('party') or ''
            v['party_type'] = v['party_type'] or doc.get('party_type') or ''
        if v['party_type'] == 'Customer':
            customers.add(v['party'])
        if v['party_type'] == 'Supplier':
            suppliers.add(v['party'])
    customer_map = _master_fields('Customer', customers - {''}, ['custom_vat_registration_number', 'tax_id'])
    supplier_map = _master_fields('Supplier', suppliers - {''}, ['custom_vat_registration_number', 'tax_id'])
    settings = _settings(filters.get('account_classification'))
    account_map = {r.get('account'): r.get('category') if r.get('category') in
                   ('Output', 'Input', 'Adjustment', 'Excluded') else 'Unclassified' for r in settings}
    results = []
    for (vtype, name), v in sorted(vouchers.items(), key=lambda x: (str(x[1]['posting_date']), x[0])):
        doc = docs.get(vtype, {}).get(name)
        section = _section(vtype, doc)
        master = (customer_map if v['party_type'] == 'Customer' else supplier_map).get(v['party'], {})
        vat = master.get('custom_vat_registration_number') or master.get('tax_id') or ''
        details = []
        for account, a in sorted(v['accounts'].items()):
            details.append({'account': account, 'debit': a['debit'], 'credit': a['credit'],
                            'tax_amount': a['credit'] - a['debit'],
                            'category': account_map.get(account) or 'Unclassified'})
        tax = sum(a['tax_amount'] for a in details)
        net = _invoice_net(vtype, doc)
        warnings = []
        if v['party_type'] in ('Customer', 'Supplier') and not vat:
            warnings.append('الرقم الضريبي غير متوفر')
        if net is None:
            warnings.append('الوعاء الضريبي غير متاح؛ مبلغ السداد ليس وعاءً ضريبيًا')
        if any(a['category'] not in ('Output', 'Input', 'Adjustment', 'Excluded') for a in details):
            warnings.append('حساب ضريبة غير مصنف')
        if v['has_tax_gl'] and not doc:
            warnings.append('المستند المصدر غير متاح؛ المبلغ مأخوذ من الأستاذ')
        if doc and doc.get('docstatus') is not None and doc.get('docstatus') != 1:
            warnings.append('حالة المستند لا تتفق مع حركة الأستاذ غير الملغاة')
        if sum(a['debit'] for a in details) and sum(a['credit'] for a in details):
            warnings.append('توجد حركات مدينة ودائنة؛ راجع تفصيل الحسابات')
        if net is not None and net * tax < 0:
            warnings.append('إشارة الضريبة تختلف عن اتجاه صافي الفاتورة')
        row = {'invoice_no': name, 'voucher_type': vtype, 'posting_date': v['posting_date'],
               'party': v['party'], 'account': ', '.join(a['account'] for a in details),
               'custom_vat_registration_number': vat, 'item_name': '', 'net_amount': net,
               'tax_amount': tax, 'section_key': section, 'indent': 1, 'row_kind': 'document',
               'account_details': details, 'warnings': warnings, 'has_tax_gl': v['has_tax_gl']}
        results.append(row)
    return results


def _reconciliation(ledger, documents, visible):
    # Independently total raw GL rows and document aggregates by account.
    result = {}
    for r in ledger:
        a = result.setdefault(r.account, {'account': r.account, 'debit': 0.0, 'credit': 0.0,
                                         'document_tax': 0.0, 'visible_tax': 0.0})
        a['debit'] += flt(r.debit)
        a['credit'] += flt(r.credit)
    for rows, field in ((documents, 'document_tax'), (visible, 'visible_tax')):
        for row in rows:
            for detail in row['account_details']:
                result[detail['account']][field] += detail['tax_amount']
    for a in result.values():
        a['ledger_tax'] = a['credit'] - a['debit']
        a['excluded_by_filters'] = a['document_tax'] - a['visible_tax']
        a['difference'] = a['ledger_tax'] - a['document_tax']
    return sorted(result.values(), key=lambda a: a['account'])


def _summary(documents):
    totals = defaultdict(float)
    for row in documents:
        for detail in row['account_details']:
            totals[detail['category']] += detail['tax_amount']
    return {'output': totals['Output'], 'input': -totals['Input'], 'adjustment': totals['Adjustment'],
            'unclassified': totals['Unclassified'], 'excluded': totals['Excluded'],
            'net': sum(row['tax_amount'] for row in documents)}


def _table_data(documents, filters):
    items = defaultdict(list)
    if cint(filters.get('show_items')):
        for vtype in ('Sales Invoice', 'Purchase Invoice'):
            names = [r['invoice_no'] for r in documents if r['voucher_type'] == vtype]
            if names:
                for item in frappe.get_all(vtype + ' Item', filters={'parent': ['in', names]},
                        fields=['parent', 'item_name'], order_by='idx asc'):
                    items[(vtype, item.parent)].append(item.item_name)
    grouped = defaultdict(list)
    for row in documents:
        grouped[row['section_key']].append(row)
    data = []
    for section in SECTION_ORDER:
        rows = grouped[section]
        if not rows:
            continue
        data.append({'invoice_no': SECTION_NAMES[section], 'row_kind': 'section', 'indent': 0})
        for row in rows:
            data.append(row)
            if cint(filters.get('show_tax_details')):
                for a in row['account_details']:
                    data.append({'voucher_type': row['voucher_type'], 'account': a['account'],
                                 'tax_amount': a['tax_amount'], 'net_amount': None,
                                 'item_name': a['category'], 'row_kind': 'tax_detail', 'indent': 2})
            for name in items[(row['voucher_type'], row['invoice_no'])]:
                data.append({'voucher_type': row['voucher_type'], 'item_name': name,
                             'row_kind': 'item', 'indent': 2, 'net_amount': None, 'tax_amount': None})
        known = [r['net_amount'] for r in rows if r['net_amount'] is not None]
        data.append({'invoice_no': _('Total') + ' ' + SECTION_NAMES[section], 'row_kind': 'subtotal',
                     'indent': 0, 'section_key': section, 'net_amount': sum(known) if known else None,
                     'tax_amount': sum(r['tax_amount'] for r in rows)})
    if documents:
        known = [r['net_amount'] for r in documents if r['net_amount'] is not None]
        data.append({'invoice_no': _('Grand Total'), 'row_kind': 'grand_total', 'indent': 0,
                     'net_amount': sum(known) if known else None,
                     'tax_amount': sum(r['tax_amount'] for r in documents)})
    return data


def _finite(value):
    if value is None or value == '':
        return None
    try:
        value = float(value)
    except (TypeError, ValueError):
        frappe.throw(_('Invalid manual amount.'))
    if not math.isfinite(value):
        frappe.throw(_('Manual amounts must be finite.'))
    return value


def _declaration(documents, filters):
    mappings = _settings(filters.get('declaration_mapping'))
    mapping = {}
    for m in mappings:
        key = (m.get('voucher_type'), m.get('invoice_no'))
        if key in mapping:
            frappe.throw(_('A voucher must be mapped only once in the declaration.'))
        mapping[key] = m
    boxes = {n: {'number': n, 'label': label, 'amount': None, 'adjustment': None, 'tax': None}
             for n, label in RETURN_LABELS.items()}
    pending = []
    for row in documents:
        m = mapping.get((row['voucher_type'], row['invoice_no']))
        number = cint(m.get('box')) if m else 0
        if number not in boxes:
            pending.append({'voucher_type': row['voucher_type'], 'invoice_no': row['invoice_no'],
                            'tax_amount': row['tax_amount'], 'reason': 'المستند غير مصنف في الإقرار'})
            continue
        box = boxes[number]
        # One explicit voucher mapping prevents repeating an invoice base per account.
        base = _finite(m.get('base_amount'))
        if base is None and cint(m.get('use_invoice_base')) and row['net_amount'] is not None:
            base = row['net_amount'] * (-1 if row['voucher_type'] == 'Purchase Invoice' else 1)
        if base is None:
            pending.append({'voucher_type': row['voucher_type'], 'invoice_no': row['invoice_no'],
                            'tax_amount': row['tax_amount'], 'reason': 'الوعاء لم يُحدد أو يُعتمد'})
        else:
            box['amount'] = flt(box['amount']) + base
        box['tax'] = flt(box['tax']) + row['tax_amount'] * (1 if number < 6 else -1)
        if number in (3, 4, 5, 10, 11) and abs(row['tax_amount']) > 0.0000001:
            pending.append({'voucher_type': row['voucher_type'], 'invoice_no': row['invoice_no'],
                            'tax_amount': row['tax_amount'], 'reason': 'ضريبة غير صفرية في بند صفري أو معفى'})
    manual_rows = _settings(filters.get('declaration_manual'))
    seen = set()
    for manual in manual_rows:
        n = cint(manual.get('box'))
        if n not in boxes or n in seen:
            frappe.throw(_('Manual declaration rows must have unique valid box numbers.'))
        seen.add(n)
        for field in ('amount', 'adjustment', 'tax'):
            value = _finite(manual.get(field))
            if value is not None:
                boxes[n][field] = flt(boxes[n][field]) + value
    result = []
    for start, end, total_number, label in ((1, 5, 6, 'إجمالي المبيعات'), (7, 11, 12, 'إجمالي المشتريات')):
        group = [boxes[n] for n in range(start, end + 1)]
        result.extend(group)
        # Missing categories stay blank in rows and make the declaration incomplete.
        result.append({'number': total_number, 'label': label, 'is_total': True,
                       **{field: sum(flt(r[field]) for r in group) for field in ('amount', 'adjustment', 'tax')}})
    missing = [n for n, r in boxes.items() if any(r[f] is None for f in ('amount', 'adjustment', 'tax'))]
    current = result[5]['tax'] - result[-1]['tax']
    previous = _finite(filters.get('previous_correction'))
    carried = _finite(filters.get('carried_credit'))
    if carried is not None and carried < 0:
        frappe.throw(_('Carried credit must be entered as a positive amount.'))
    return {'rows': result, 'pending': pending, 'missing_boxes': missing, 'current_tax': current,
            'previous_correction': previous, 'carried_credit': carried,
            'payable': current + (previous or 0) - (carried or 0),
            'complete': not pending and not missing and previous is not None and carried is not None,
            'manual_entries': manual_rows}


def _build(filters):
    filters = _validate(frappe._dict(_json(filters, {})))
    ledger = _read_ledger(filters)
    documents = _load_documents(filters, ledger)
    visible = [r for r in documents if _visible(r, filters)]
    warnings = [{'voucher_type': r['voucher_type'], 'invoice_no': r['invoice_no'], 'message': w}
                for r in visible for w in r['warnings']]
    return {'filters': dict(filters), 'generated_at': str(now_datetime()),
            'currency': frappe.db.get_value('Company', filters.company, 'default_currency'),
            'documents': visible, 'data': _table_data(visible, filters),
            'reconciliation': _reconciliation(ledger, documents, visible),
            'summary': _summary(visible), 'warnings': warnings,
            'reconciliation_details': [{'voucher_type': r['voucher_type'], 'invoice_no': r['invoice_no'],
                                       **a} for r in documents for a in r['account_details']]}


def get_columns():
    columns = [
        {
            "label": _("Voucher No"),
            "fieldname": "invoice_no",
            "fieldtype": "Dynamic Link",
            "options": "voucher_type",
            "width": 250,
        },
        {
            "label": _("Voucher Type"),
            "fieldname": "voucher_type",
            "fieldtype": "Data",
            "width": 150,
        },
        {
            "label": _("Party"),
            "fieldname": "party",
            "fieldtype": "Data",
            "width": 150,
        },
        {
            "label": _("VAT Registration Number"),
            "fieldname": "custom_vat_registration_number",
            "fieldtype": "Data",
            "width": 180,
        },
        {
            "label": _("Date"),
            "fieldname": "posting_date",
            "fieldtype": "Date",
            "width": 120,
        },
        {
            "label": _("Tax Account"),
            "fieldname": "account",
            "fieldtype": "Data",
            "width": 200,
        },
        {
            "label": _("Item"),
            "fieldname": "item_name",
            "fieldtype": "Data",
            "width": 200,
        },
        {
            "label": _("Net Amount"),
            "fieldname": "net_amount",
            "fieldtype": "Currency",
            "width": 150,
        },
        {
            "label": _("Tax Amount"),
            "fieldname": "tax_amount",
            "fieldtype": "Currency",
            "width": 150,
        },
    ]


    return columns


def execute(filters=None):
    result = _build(filters or {})
    s = result['summary']
    labels = [('output', 'ضريبة المخرجات', 'Green'), ('input', 'ضريبة المدخلات', 'Blue'),
              ('adjustment', 'التسويات المصنفة', 'Orange'), ('unclassified', 'غير مصنف (صافي)', 'Orange'),
              ('excluded', 'مستبعد من الإقرار (صافي)', 'Gray'), ('net', 'صافي حركة الضريبة', 'Blue')]
    summary = [{'label': label, 'value': s[key], 'datatype': 'Currency',
                'currency': result['currency'], 'indicator': color} for key, label, color in labels]
    message = 'المبالغ بعملة الشركة. صافي الفاتورة كامل وليس موزعًا على الحسابات المختارة. '
    message += 'الخانة الفارغة تعني وعاءً غير متاح؛ المجاميع تشمل الصافي المعروف فقط. '
    if cint(result['filters'].get('show_tax_details')):
        message += 'تفاصيل الحسابات للشرح؛ إجمالي المستند يُحتسب مرة واحدة فقط. '
    message += 'تنبيهات المراجعة: ' + str(len(result['warnings']))
    return get_columns(), result['data'], message, None, summary, 1


@frappe.whitelist()
def get_print_data(filters=None, layout='summary'):
    result = _build(filters or {})
    result['columns'] = get_columns()
    result['layout'] = layout if layout in ('summary', 'detailed', 'declaration') else 'summary'
    if result['layout'] == 'declaration':
        result['declaration'] = _declaration(result['documents'], frappe._dict(result['filters']))
        restricted = any(result['filters'].get(k) for k in ('party', 'invoice_no', 'voucher_type'))
        restricted = restricted or result['filters'].get('document_group', 'All') != 'All'
        restricted = restricted or not cint(result['filters'].get('include_zero_tax'))
        restricted = restricted or not cint(result['filters'].get('include_non_taxed'))
        result['declaration']['restricted_scope'] = bool(restricted)
        if restricted or result['currency'] != 'SAR':
            result['declaration']['complete'] = False
    return result


@frappe.whitelist()
def get_review_data(filters=None):
    return _build(filters or {})


@frappe.whitelist()
def export_excel(filters=None):
    from io import BytesIO
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    report = get_print_data(filters, 'declaration')
    wb = Workbook()
    wb.remove(wb.active)
    def sheet(name, headers, rows):
        ws = wb.create_sheet(name)
        ws.sheet_view.rightToLeft = True
        ws.append(headers)
        for row in rows:
            ws.append(row)
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.fill = PatternFill('solid', fgColor='087F7B')
            cell.font = Font(color='FFFFFF', bold=True)
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                # Names and notes must never become Excel formulas.
                if isinstance(cell.value, str):
                    cell.data_type = 's'
                elif isinstance(cell.value, (int, float)):
                    cell.number_format = '#,##0.00;[Red](#,##0.00)'
                cell.alignment = Alignment(vertical='top', wrap_text=True)
        for i in range(1, len(headers) + 1):
            ws.column_dimensions[get_column_letter(i)].width = 24
        return ws
    columns = report['columns']
    sheet('التقرير', [c['label'] for c in columns] + ['نوع الصف', 'تنبيهات'],
          [[r.get(c['fieldname']) for c in columns] + [r.get('row_kind'), '؛ '.join(r.get('warnings', []))]
           for r in report['data']])
    sheet('تفصيل الحسابات', ['نوع المستند', 'المستند', 'الحساب', 'مدين', 'دائن', 'صافي الضريبة', 'التصنيف'],
          [[r['voucher_type'], r['invoice_no'], a['account'], a['debit'], a['credit'], a['tax_amount'], a['category']]
           for r in report['documents'] for a in r['account_details']])
    fields = ['account', 'debit', 'credit', 'ledger_tax', 'document_tax', 'visible_tax', 'excluded_by_filters', 'difference']
    sheet('المطابقة', ['الحساب', 'مدين', 'دائن', 'صافي الأستاذ', 'صافي المستندات', 'المعروض', 'مستبعد بالفلاتر', 'الفرق'],
          [[r[f] for f in fields] for r in report['reconciliation']])
    d = report['declaration']
    sheet('الإقرار', ['البند', 'البيان', 'المبلغ', 'التعديل', 'الضريبة'],
          [[r['number'], r['label'], r['amount'], r['adjustment'], r['tax']] for r in d['rows']] +
          [[None, 'صافي الضريبة للفترة', None, None, d['current_tax']],
           [None, 'تصحيحات الفترات السابقة', None, None, d['previous_correction']],
           [None, 'الرصيد المرحّل', None, None, d['carried_credit']],
           [None, 'صافي مستحق مبدئي', None, None, d['payable']]])
    sheet('المراجعة', ['نوع المستند', 'المستند', 'الملاحظة'],
          [[r['voucher_type'], r['invoice_no'], r['message']] for r in report['warnings']] +
          [[r['voucher_type'], r['invoice_no'], r['reason']] for r in d['pending']])
    meta = [['الشركة', report['filters']['company']], ['العملة', report['currency']],
            ['وقت الاستخراج', report['generated_at']], ['حالة الإقرار', 'مكتمل التصنيف' if d['complete'] else 'مسودة غير مكتملة'],
            ['ملاحظة', 'صافي الفواتير كامل؛ المجاميع تشمل الأوعية المعروفة فقط. تفاصيل الحسابات لا تُجمع مرة ثانية.'],
            ['مرجع شكل الإقرار', 'https://zatca.gov.sa/ar/HelpCenter/guidelines/Documents/إرشادات.pdf']]
    meta += [[k, json.dumps(v, ensure_ascii=False, default=str)] for k, v in report['filters'].items()]
    sheet('الإعدادات', ['البيان', 'القيمة'], meta)
    stream = BytesIO()
    wb.save(stream)
    frappe.local.response.filename = 'All-Tax-Report.xlsx'
    frappe.local.response.filecontent = stream.getvalue()
    frappe.local.response.type = 'binary'
